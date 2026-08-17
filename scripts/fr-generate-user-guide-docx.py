"""Generate the local bilingual Word guide from its maintained Markdown source."""

from __future__ import annotations

import re
import sys
import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "USER_GUIDE.zh-CN.en.md"
OUTPUT = ROOT / "dist" / "docs" / "FR-ComfyUI-ControlCenter-User-Guide.zh-CN.en.docx"
APP_VERSION = json.loads((ROOT / "package.json").read_text(encoding="utf-8"))["version"]


def add_runs(paragraph, text: str) -> None:
    """Render the small inline Markdown subset used by the guide."""

    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\[.+?\]\(.+?\))")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Cascadia Mono"
            run.font.size = Pt(9)
        else:
            label, url = token[1:-1].split("](", 1)
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor(5, 99, 193)
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_table(document: Document, rows: list[str]) -> None:
    cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
    data = [row for row in cells if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)]
    if not data:
        return

    table = document.add_table(rows=1, cols=len(data[0]))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for index, value in enumerate(data[0]):
        header[index].text = value
        for run in header[index].paragraphs[0].runs:
            run.bold = True
    for row in data[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            if index < len(cells):
                add_runs(cells[index].paragraphs[0], value)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def build_document(source: Path, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)

    code_style = document.styles.add_style("FR Code", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Cascadia Mono"
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.left_indent = Cm(0.6)
    code_style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "FR ComfyUI Control Center | User Guide"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"FR ComfyUI Control Center | {APP_VERSION} | ")
    footer.add_run("Page ")
    add_page_number(footer)

    lines = source.read_text(encoding="utf-8").splitlines()
    table_rows: list[str] = []
    in_code = False

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(document, table_rows)
            table_rows = []

    for raw_line in lines:
        line = raw_line.rstrip()
        if line.startswith("|"):
            table_rows.append(line)
            continue
        flush_table()

        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            document.add_paragraph(line, style="FR Code")
            continue
        if not line or line == "---":
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 4)
            text = line[level:].strip()
            paragraph = document.add_heading(text, level=level)
            if level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.7)
            paragraph.paragraph_format.right_indent = Cm(0.3)
            add_runs(paragraph, line[2:])
            for run in paragraph.runs:
                run.italic = True
            continue
        if line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, line[2:])
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            paragraph = document.add_paragraph(style="List Number")
            add_runs(paragraph, numbered.group(2))
            continue
        paragraph = document.add_paragraph()
        add_runs(paragraph, line)

    flush_table()
    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "FR ComfyUI Control Center User Guide"
    document.core_properties.subject = "Bilingual local operating guide"
    document.core_properties.author = "FR AI"
    document.save(output)


if __name__ == "__main__":
    if not SOURCE.is_file():
        raise SystemExit(f"Guide source is missing: {SOURCE}")
    build_document(SOURCE, OUTPUT)
    print(OUTPUT)
